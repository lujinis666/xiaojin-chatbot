import os
import io
import re as _re
import sqlite3
import json
from flask import Flask, request, jsonify, render_template, send_file, g, Response, stream_with_context
from dotenv import load_dotenv
from langchain.chat_models import init_chat_model
from langchain_core.messages import HumanMessage, SystemMessage, AIMessage

load_dotenv()

app = Flask(__name__)

# ── 数据库路径 ────────────────────────────────────────────────
BASE_DIR = os.path.dirname(os.path.abspath(__file__))
DB_PATH  = os.path.join(BASE_DIR, '../data', 'history.db')

# ── 模型初始化 ───────────────────────────────────────────────
base_url = os.getenv("LONGCAT_BASE_URL")
api_key  = os.getenv("LONGCAT_API_KEY")

model = init_chat_model(
    model="LongCat-2.0-Preview",
    model_provider="openai",
    base_url=base_url,
    api_key=api_key,
    timeout=90,       # 请求超时90秒则报错，避免永久挂起
    max_retries=0,    # 不重试，失败就失败
)

# ══════════════════════════════════════════════════════════════
#  SQLite 数据库层
# ══════════════════════════════════════════════════════════════

def get_db() -> sqlite3.Connection:
    """获取当前请求的数据库连接（通过 Flask g 缓存）"""
    if '_db' not in g:
        os.makedirs(os.path.dirname(DB_PATH), exist_ok=True)
        conn = sqlite3.connect(DB_PATH, timeout=15)  # 等待锁最多 15 秒
        conn.row_factory = sqlite3.Row
        conn.execute("PRAGMA journal_mode=WAL")
        conn.execute("PRAGMA foreign_keys=ON")
        conn.execute("PRAGMA synchronous=NORMAL")  # 比 FULL 快，仍保证数据安全
        g._db = conn
    return g._db

@app.teardown_appcontext
def close_db(e=None):
    db = g.pop('_db', None)
    if db:
        db.close()

def init_db():
    """创建数据库表（幂等操作）"""
    os.makedirs(os.path.dirname(DB_PATH), exist_ok=True)
    with sqlite3.connect(DB_PATH) as conn:
        conn.executescript("""
            CREATE TABLE IF NOT EXISTS sessions (
                session_id TEXT     PRIMARY KEY,
                title      TEXT     NOT NULL DEFAULT '新对话',
                created_at DATETIME NOT NULL DEFAULT CURRENT_TIMESTAMP,
                updated_at DATETIME NOT NULL DEFAULT CURRENT_TIMESTAMP
            );

            CREATE TABLE IF NOT EXISTS messages (
                id         INTEGER  PRIMARY KEY AUTOINCREMENT,
                session_id TEXT     NOT NULL,
                role       TEXT     NOT NULL,   -- 'user' | 'assistant'
                content    TEXT     NOT NULL,
                created_at DATETIME NOT NULL DEFAULT CURRENT_TIMESTAMP,
                FOREIGN KEY (session_id) REFERENCES sessions(session_id) ON DELETE CASCADE
            );

            CREATE INDEX IF NOT EXISTS idx_messages_session
                ON messages(session_id, created_at);
        """)

init_db()

# ── 数据库辅助函数 ───────────────────────────────────────────

def load_recent_history(db: sqlite3.Connection, session_id: str, limit: int = 20):
    """加载最近 limit 条消息（按时间正序返回）"""
    rows = db.execute(
        """SELECT role, content FROM messages
           WHERE session_id=?
           ORDER BY created_at DESC LIMIT ?""",
        (session_id, limit)
    ).fetchall()
    return list(reversed(rows))   # 反转为时间正序


def get_message_count(db: sqlite3.Connection, session_id: str) -> int:
    row = db.execute(
        "SELECT COUNT(*) FROM messages WHERE session_id=?", (session_id,)
    ).fetchone()
    return row[0] if row else 0


def setup_session_and_save(
    db: sqlite3.Connection,
    session_id: str,
    user_msg: str,
    reply: str
):
    """
    在单次事务中完成：
    1. 确保 session 存在
    2. 首条消息设为标题
    3. 保存 user + assistant 消息
    4. 更新 session 时间戳
    """
    # 读操作（事务外）
    count = get_message_count(db, session_id)

    # 单次写事务
    title = user_msg[:30].strip()
    with db:  # with db 会自动 commit 或 rollback
        db.execute(
            "INSERT OR IGNORE INTO sessions (session_id, title) VALUES (?, ?)",
            (session_id, title)
        )
        if count == 0:
            db.execute(
                "UPDATE sessions SET title=? WHERE session_id=?",
                (title, session_id)
            )
        db.execute(
            "INSERT INTO messages (session_id, role, content) VALUES (?, 'user', ?)",
            (session_id, user_msg)
        )
        db.execute(
            "INSERT INTO messages (session_id, role, content) VALUES (?, 'assistant', ?)",
            (session_id, reply)
        )
        db.execute(
            "UPDATE sessions SET updated_at=CURRENT_TIMESTAMP WHERE session_id=?",
            (session_id,)
        )


# ══════════════════════════════════════════════════════════════
#  DuckDuckGo 搜索
# ══════════════════════════════════════════════════════════════

def do_search(query: str, max_results: int = 5) -> str:
    try:
        from ddgs import DDGS
        with DDGS() as ddgs:
            results = list(ddgs.text(query, max_results=max_results))
        if not results:
            return ""
        parts = []
        for i, r in enumerate(results, 1):
            parts.append(
                f"[{i}] {r.get('title','')}\n{r.get('body','')}\n来源：{r.get('href','')}"
            )
        return "\n\n".join(parts)
    except Exception as e:
        print(f"[Search Error] {e}")
        return ""


# ══════════════════════════════════════════════════════════════
#  Markdown → Word 文档
# ══════════════════════════════════════════════════════════════

def markdown_to_docx(content: str) -> io.BytesIO:
    from docx import Document
    from docx.shared import Pt, RGBColor, Inches
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Arial"
    style.font.size = Pt(11)

    def add_inline(para, text: str):
        tokens = _re.split(r'(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`)', text)
        for t in tokens:
            if t.startswith("**") and t.endswith("**"):
                doc_run = para.add_run(t[2:-2]); doc_run.bold = True
            elif t.startswith("*") and t.endswith("*"):
                doc_run = para.add_run(t[1:-1]); doc_run.italic = True
            elif t.startswith("`") and t.endswith("`"):
                doc_run = para.add_run(t[1:-1])
                doc_run.font.name = "Courier New"
                doc_run.font.color.rgb = RGBColor(0xD9, 0x48, 0x00)
            else:
                para.add_run(t)

    lines = content.split("\n")
    in_code, code_lines = False, []
    for line in lines:
        if line.strip().startswith("```"):
            if not in_code:
                in_code = True; code_lines = []
            else:
                in_code = False
                for cl in code_lines:
                    p = doc.add_paragraph(style="No Spacing")
                    run = p.add_run(cl)
                    run.font.name = "Courier New"; run.font.size = Pt(9.5)
                    run.font.color.rgb = RGBColor(0x2D, 0x74, 0xDA)
                    p.paragraph_format.left_indent = Inches(0.3)
            continue
        if in_code:
            code_lines.append(line); continue

        if   line.startswith("#### "): doc.add_heading(line[5:], level=4)
        elif line.startswith("### "):  doc.add_heading(line[4:], level=3)
        elif line.startswith("## "):   doc.add_heading(line[3:], level=2)
        elif line.startswith("# "):    doc.add_heading(line[2:], level=1)
        elif line.startswith("- ") or line.startswith("* "):
            p = doc.add_paragraph(style="List Bullet"); add_inline(p, line[2:])
        elif _re.match(r"^\d+\. ", line):
            p = doc.add_paragraph(style="List Number")
            add_inline(p, _re.sub(r"^\d+\. ", "", line))
        elif line.strip() in ("---", "***", "___"):
            p = doc.add_paragraph()
            pPr = p._p.get_or_add_pPr()
            pBdr = OxmlElement("w:pBdr")
            bottom = OxmlElement("w:bottom")
            bottom.set(qn("w:val"), "single"); bottom.set(qn("w:sz"), "6")
            bottom.set(qn("w:space"), "1"); bottom.set(qn("w:color"), "CCCCCC")
            pBdr.append(bottom); pPr.append(pBdr)
        elif line.strip() == "":
            doc.add_paragraph("")
        else:
            p = doc.add_paragraph(); add_inline(p, line)

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf


# ══════════════════════════════════════════════════════════════
#  路由
# ══════════════════════════════════════════════════════════════

@app.route("/")
def home():
    return render_template("index.html")


# ── 会话管理 API ─────────────────────────────────────────────

@app.route("/api/sessions", methods=["GET"])
def list_sessions():
    db = get_db()
    rows = db.execute(
        """SELECT s.session_id, s.title, s.created_at, s.updated_at,
                  COUNT(m.id) AS message_count
           FROM sessions s
           LEFT JOIN messages m ON s.session_id = m.session_id
           GROUP BY s.session_id
           ORDER BY s.updated_at DESC
           LIMIT 100"""
    ).fetchall()
    return jsonify([dict(r) for r in rows])


@app.route("/api/sessions/<session_id>", methods=["PUT"])
def rename_session(session_id):
    data  = request.json or {}
    title = (data.get("title") or "").strip()
    if not title:
        return jsonify({"error": "标题不能为空"}), 400
    db = get_db()
    with db:
        db.execute(
            "UPDATE sessions SET title=? WHERE session_id=?",
            (title[:50], session_id)
        )
    return jsonify({"ok": True})


@app.route("/api/sessions/<session_id>", methods=["DELETE"])
def delete_session(session_id):
    db = get_db()
    db.execute("DELETE FROM messages WHERE session_id=?", (session_id,))
    db.execute("DELETE FROM sessions WHERE session_id=?", (session_id,))
    db.commit()
    return jsonify({"ok": True})


@app.route("/api/sessions/<session_id>/messages", methods=["GET"])
def get_session_messages(session_id):
    db = get_db()
    rows = db.execute(
        """SELECT role, content, created_at
           FROM messages WHERE session_id=?
           ORDER BY created_at ASC""",
        (session_id,)
    ).fetchall()
    return jsonify([dict(r) for r in rows])


# ── 聊天 API ─────────────────────────────────────────────────

@app.route("/api/chat", methods=["POST"])
def chat():
    data           = request.json or {}
    user_msg       = data.get("message", "").strip()
    session_id     = data.get("session_id", "default")
    search_enabled = data.get("search_enabled", False)

    if not user_msg:
        return jsonify({"error": "消息不能为空"}), 400

    db = get_db()

    try:
        # ── Step 1：联网搜索 ───────────────────────────────────
        tool_calls_info = []
        search_context  = ""

        if search_enabled:
            print(f"[Search] 搜索：{user_msg}")
            result = do_search(user_msg)
            if result:
                search_context = result
                tool_calls_info.append({"tool": "web_search", "input": user_msg})

        # ── Step 2：构建系统提示 ────────────────────────────────
        system_content = (
            "你是 LongCat，一个由 LongCat 团队开发的智能对话助手。"
            "你博学友善，擅长中文交流，回答清晰专业。"
        )
        if search_context:
            system_content += (
                "\n\n以下是通过 DuckDuckGo 网络搜索获取的最新资讯，"
                "请优先基于这些内容回答用户问题，并在回答末尾注明来源于网络搜索：\n\n"
                + search_context
            )

        # ── Step 3：加载历史 + 构建消息列表 ─────────────────────
        history = load_recent_history(db, session_id, limit=20)
        messages = [SystemMessage(content=system_content)]
        for h in history:
            if h["role"] == "user":
                messages.append(HumanMessage(content=h["content"]))
            else:
                messages.append(AIMessage(content=h["content"]))
        messages.append(HumanMessage(content=user_msg))

        # ── Step 4 & 5：流式调用模型并返回 SSE ─────────────────
        def generate():
            try:
                if search_enabled and search_context:
                    # 推送工具调用状态
                    yield f"data: {json.dumps({'type': 'tool', 'tool': 'web_search', 'input': user_msg})}\n\n"

                full_reply = []
                for chunk in model.stream(messages):
                    if chunk.content:
                        full_reply.append(chunk.content)
                        # 推送生成的 token
                        yield f"data: {json.dumps({'type': 'chunk', 'content': chunk.content})}\n\n"

                # 只有在完全生成后才持久化
                reply_str = "".join(full_reply)
                
                # 使用独立的 db 事务保存对话，避免因为流式生成过长导致 g.db 被 Flask 过早关闭
                with sqlite3.connect(DB_PATH, timeout=15) as conn:
                    conn.row_factory = sqlite3.Row
                    conn.execute("PRAGMA journal_mode=WAL")
                    setup_session_and_save(conn, session_id, user_msg, reply_str)

                # 标记结束
                yield f"data: {json.dumps({'type': 'done'})}\n\n"

            except TimeoutError:
                print("[Chat Timeout] 模型响应超时")
                yield f"data: {json.dumps({'type': 'error', 'error': '模型响应超时，请稍后重试'})}\n\n"
            except Exception as e:
                print(f"[Chat Error] {type(e).__name__}: {e}")
                yield f"data: {json.dumps({'type': 'error', 'error': f'调用失败：{str(e)}'})}\n\n"

        return Response(stream_with_context(generate()), mimetype="text/event-stream")

    except Exception as e:
        print(f"[Chat Setup Error] {type(e).__name__}: {e}")
        return jsonify({"error": f"请求初始化失败：{str(e)}"}), 500


# ── 文档导出 ─────────────────────────────────────────────────

@app.route("/api/export-docx", methods=["POST"])
def export_docx():
    try:
        import docx  # noqa
    except ImportError:
        return jsonify({"error": "请先安装 python-docx：uv add python-docx"}), 500

    data     = request.json or {}
    content  = data.get("content", "")
    filename = (data.get("filename", "") or "LongCat文档").strip()

    if not content:
        return jsonify({"error": "内容不能为空"}), 400

    try:
        buf = markdown_to_docx(content)
        return send_file(
            buf,
            mimetype="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            as_attachment=True,
            download_name=f"{filename}.docx",
        )
    except Exception as e:
        return jsonify({"error": f"文档生成失败：{str(e)}"}), 500


if __name__ == "__main__":
    app.run(debug=True, port=5000)